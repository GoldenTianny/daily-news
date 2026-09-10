#!/usr/bin/env python3
"""수정주가 전체 이력 파일(Peer Analysis 배열) -> 가격 DB 전면 갱신 + 종목명 일괄 정정

원본: HTS 'Peer Analysis (Time Series)' 다운로드. 행 = 종목(Code, Name, 결산월),
열 = 일자(Period 행 YYYYMMDD), 값 = 수정주가. 액면분할·합병 등으로 과거치가 소급
조정되거나 사명이 바뀐 종목을 한 번에 반영하기 위한 용도 (일일 스냅샷은 당일만 갱신하므로).

처리
1) 가격: (date, code) 단위 upsert — 원본에 있는 종목·날짜만 교체, ETF 등 원본에 없는
   코드의 기존 행은 유지. 원본이 DB보다 과거(2023-12-28~)를 담고 있으면 이력이 늘어남
2) 종목명: 코드별로 원본의 현재 사명('(주)'·'㈜' 접두 제거)을 정본으로 삼아
   가격·목표주가·RS·미너비니·영업이익(annual, consensus) DB의 name을 코드 기준으로 일괄 교체
   (검색기는 종목명으로 DB를 조인하므로 사명변경 종목의 차트·실적이 끊기지 않게)
3) 보고: 과거치가 바뀐 종목 수(분할·합병 추정), 사명변경 목록, 신규·누락 코드

이후 RS 등급 전체 재계산(build_rs --all)과 파생 스터디 재생성이 필요 — ingest_daily.py가
이 파일을 자동 인식해 순서대로 실행한다.
실행: python3 tools/market/refresh_prices.py <수정주가_YYYYMMDD.xlsx> [--dry-run]
"""
import sys, os, re, glob, numbers, datetime
import openpyxl
import duckdb
import pandas as pd

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DB = os.path.join(REPO, 'db', 'market')
PRICE_DIR = os.path.join(DB, 'price')
NAME_DBS = [   # name 컬럼을 코드 기준으로 정정할 DB (glob)
    os.path.join(DB, 'consensus', '*.parquet'),
    os.path.join(DB, 'rs', '*.parquet'),
    os.path.join(DB, 'minervini', '*.parquet'),
    os.path.join(DB, 'earnings', 'annual.parquet'),
    os.path.join(DB, 'earnings', 'consensus', '*.parquet'),
]


def is_price_history_file(xlsx_path):
    """첫 시트 상단에 'Peer Analysis' 태그 + Code/Name 헤더 행 아래 '수정주가' 항목이면 해당"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    peer = False
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 14:
            break
        first = str(row[0] or '') if row else ''
        if 'Peer Analysis' in first:
            peer = True
        if first.strip() == 'Code' and peer:
            return any(str(v or '').strip() == '수정주가' for v in row[3:8])
    return False


def clean_name(s):
    return re.sub(r'^\(주\)|^㈜', '', str(s or '').strip()).strip()


def parse(xlsx_path):
    """-> DataFrame(date, code, name, close)"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = {str(r[1] or '').strip(): i for i, r in enumerate(rows[:14]) if r and r[1]}
    period = rows[hdr['Period']]
    code_hdr = next(i for i, r in enumerate(rows[:16]) if r and str(r[0] or '').strip() == 'Code')
    item_row = rows[code_hdr]
    date_cols = []
    for c in range(3, len(period)):
        s = str(period[c] or '').strip()
        if re.fullmatch(r'\d{8}', s) and str(item_row[c] or '').strip() == '수정주가':
            date_cols.append((c, datetime.date(int(s[:4]), int(s[4:6]), int(s[6:]))))
    out = []
    for r in rows[code_hdr + 1:]:
        code = str(r[0] or '').strip()
        if not code.startswith('A'):
            continue
        nm = clean_name(r[1])
        for c, d in date_cols:
            v = r[c] if c < len(r) else None
            if isinstance(v, numbers.Number) and v > 0:
                out.append((d, code, nm, float(v)))
    return pd.DataFrame(out, columns=['date', 'code', 'name', 'close'])


def load_price_db():
    files = sorted(glob.glob(os.path.join(PRICE_DIR, '*.parquet')))
    df = pd.concat([pd.read_parquet(f) for f in files], ignore_index=True)
    df['date'] = pd.to_datetime(df['date']).dt.date
    return df


def write_monthly(df):
    """월별 파일 저장 (내용이 같은 달은 건너뜀)"""
    con = duckdb.connect()
    n_new = n_same = 0
    for ym, g in df.groupby(df['date'].map(lambda d: f'{d.year:04d}-{d.month:02d}')):
        dst = os.path.join(PRICE_DIR, ym + '.parquet')
        g = g.sort_values(['date', 'code']).reset_index(drop=True)
        if os.path.exists(dst):
            old = pd.read_parquet(dst)
            old['date'] = pd.to_datetime(old['date']).dt.date
            old = old.sort_values(['date', 'code']).reset_index(drop=True)
            if old.equals(g):
                n_same += 1
                continue
        con.execute(f"""
            COPY (SELECT CAST(date AS DATE) AS date, code, name, CAST(close AS DOUBLE) AS close FROM g)
            TO '{dst}' (FORMAT PARQUET, COMPRESSION SNAPPY)""")
        n_new += 1
    return n_new, n_same


def rename_in_db(pattern, ren):
    """glob의 각 parquet에서 code -> new name 적용 (변경된 파일만 재기록). -> 바뀐 행 수"""
    con = duckdb.connect()
    total = 0
    for f in sorted(glob.glob(pattern)):
        df = pd.read_parquet(f)
        if 'code' not in df.columns or 'name' not in df.columns:
            continue
        new = df['code'].map(ren)
        mask = new.notna() & (new != df['name'])
        if not mask.any():
            continue
        df.loc[mask, 'name'] = new[mask]
        total += int(mask.sum())
        # 스키마 보존을 위해 duckdb로 원본 타입 유지 재기록
        types = con.execute(f"DESCRIBE SELECT * FROM '{f}'").fetchall()
        cast = ', '.join(f'CAST("{c}" AS {t}) AS "{c}"' for c, t, *_ in types)
        con.execute(f"COPY (SELECT {cast} FROM df) TO '{f}' (FORMAT PARQUET, COMPRESSION SNAPPY)")
    return total


def refresh(xlsx_path, dry_run=False):
    src = parse(xlsx_path)
    if src.empty:
        print('SKIP: 수정주가 데이터 없음')
        return
    old = load_price_db()
    print(f"원본: 종목 {src['code'].nunique():,}개, {src['date'].min()} ~ {src['date'].max()}, {len(src):,}행")
    print(f"DB  : 종목 {old['code'].nunique():,}개, {old['date'].min()} ~ {old['date'].max()}, {len(old):,}행")

    # --- 비교: 가격 소급 변경(분할·합병 추정), 사명변경, 신규·누락 코드
    m = old.merge(src, on=['date', 'code'], suffixes=('_db', '_src'))
    diff = m[(m['close_db'] - m['close_src']).abs() > m['close_db'] * 1e-6]
    chg_codes = diff.groupby('code').agg(n=('date', 'size'), name=('name_src', 'first'),
                                         ratio=('close_src', lambda s: None))
    ratios = (diff['close_src'] / diff['close_db']).groupby(diff['code']).median().round(4)
    chg_codes['ratio'] = ratios
    chg_codes = chg_codes.sort_values('n', ascending=False)

    db_names = old.sort_values('date').groupby('code')['name'].last()
    src_names = src.groupby('code')['name'].first()
    common = src_names.index.intersection(db_names.index)
    renames = {c: (db_names[c], src_names[c]) for c in common if db_names[c] != src_names[c]}
    new_codes = sorted(set(src_names.index) - set(db_names.index))
    missing = sorted(set(db_names.index) - set(src_names.index))
    etf_codes = set()
    for f in glob.glob(os.path.join(REPO, 'db', 'etf', '*.parquet')):
        etf_codes |= set(pd.read_parquet(f, columns=['etf_code'])['etf_code'].unique())
    missing_non_etf = [c for c in missing if c not in etf_codes]

    print(f"\n[가격 소급 변경] 값이 달라진 종목 {len(chg_codes):,}개 / {len(diff):,}행 "
          f"(비교 가능 {len(m):,}행 중)")
    for code, r in chg_codes.head(15).iterrows():
        print(f"   {code} {r['name']:<14s} {int(r['n']):>4}일 변경 · 원본/DB 중앙값 비율 {r['ratio']}")
    print(f"[사명변경] {len(renames)}개")
    for c, (a, b) in sorted(renames.items()):
        print(f"   {c} {a} -> {b}")
    print(f"[신규 코드] DB에 없던 종목 {len(new_codes):,}개 (원본에만 존재 — 추가됨)")
    print(f"[누락 코드] 원본에 없는 DB 종목 {len(missing):,}개 (ETF {len(missing) - len(missing_non_etf):,} · "
          f"그 외 {len(missing_non_etf):,} — 기존 행 유지)")
    if dry_run:
        print('\n(dry-run: 저장하지 않음)')
        return

    # --- 가격 upsert + 종목명 정본화
    ren = {c: b for c, (_, b) in renames.items()}
    keys = set(zip(src['date'], src['code']))
    keep = old[[k not in keys for k in zip(old['date'], old['code'])]].copy()
    keep['name'] = keep['code'].map(ren).fillna(keep['name'])
    merged = pd.concat([keep, src], ignore_index=True)
    n_new, n_same = write_monthly(merged)
    print(f"\nOK  가격 DB: {len(merged):,}행, 종목 {merged['code'].nunique():,}개, "
          f"{merged['date'].min()} ~ {merged['date'].max()} | 월 파일 {n_new}개 갱신, {n_same}개 동일")

    if ren:
        for pat in NAME_DBS:
            n = rename_in_db(pat, ren)
            if n:
                print(f"OK  종목명 정정 {os.path.relpath(pat, REPO)}: {n:,}행")
    return {'renames': renames, 'changed': chg_codes, 'new': new_codes, 'missing': missing}


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit('사용법: python3 tools/market/refresh_prices.py <수정주가_YYYYMMDD.xlsx> [--dry-run]')
    refresh(sys.argv[1], dry_run='--dry-run' in sys.argv[1:])
