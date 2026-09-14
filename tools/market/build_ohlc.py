#!/usr/bin/env python3
"""수정 시가·고가·저가 시계열(Peer Analysis 배열) -> OHLC DB (db/market/ohlc/)

원본: HTS Peer Analysis 다운로드. 항목별로 시트가 나뉘며(수정시가/수정고가/수정저가),
행 = 종목(Code, Name, 결산월), 열 = 일자(Period 행 YYYYMMDD). 시트명과 무관하게
Code 헤더 행의 항목명으로 어떤 값인지 판별한다.

일일 파일(ETF_price_concensus_*.xlsx)의 스냅샷 시트에 든 기준일 1일치 수정시가·고가·저가
열(Period 'CPD')도 build_snapshot()으로 같은 DB에 병합한다 (ingest_daily 3단계).

종가(close)는 이미 db/market/price/에 있으므로 **여기엔 open·high·low만** 저장한다.
가격 DB는 검색기가 매 화면 13개월치를 읽으므로 컬럼을 늘리지 않고 분리 (백테스트 전용).

산출물: db/market/ohlc/YYYY-MM.parquet
  date DATE, code VARCHAR, name VARCHAR, open/high/low DOUBLE (원)
- 원본에 있는 (date, code)만 교체하는 upsert. 내용이 같은 달은 다시 쓰지 않음
- 적재 후 reconcile_price()로 가격 DB와 정합성 점검: OHLC 파일이 수정주가 파일보다 최신이면
  그 사이 시행된 액면분할·병합이 고가·저가에만 반영돼 close가 [low, high] 밖으로 나간다.
  close × r 이 거의 모든 날 범위 안에 들어가는 배율 r을 찾아 price DB를 소급 보정
- 종가와 함께 쓰려면 price DB와 조인:
    SELECT o.*, p.close FROM 'db/market/ohlc/*.parquet' o
    JOIN 'db/market/price/*.parquet' p USING (date, code)
실행: python3 tools/market/build_ohlc.py <파일.xlsx>   (ingest_daily.py가 시트 구조로 자동 인식)
"""
import sys, os, re, glob, numbers, datetime
import openpyxl
import duckdb
import pandas as pd

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(REPO, 'db', 'market', 'ohlc')
ITEMS = {'수정시가': 'open', '수정고가': 'high', '수정저가': 'low'}
VALS = ['open', 'high', 'low']


def clean_name(s):
    return re.sub(r'^\(주\)|^㈜', '', str(s or '').strip()).strip()


def item_sheets(wb):
    """-> [(시트명, 필드)] — Code 헤더 행의 항목명이 수정시가/고가/저가인 시트"""
    out = []
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 16:
                break
            if row and str(row[0] or '').strip() == 'Code':
                for v in row[3:8]:
                    f = ITEMS.get(str(v or '').strip())
                    if f:
                        out.append((ws.title, f))
                        break
                break
    return out


def is_ohlc_file(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    return bool(item_sheets(wb))


def parse_sheet(ws, field):
    rows = list(ws.iter_rows(values_only=True))
    hdr = {str(r[1] or '').strip(): i for i, r in enumerate(rows[:16]) if r and r[1]}
    code_hdr = next(i for i, r in enumerate(rows[:18]) if r and str(r[0] or '').strip() == 'Code')
    period, item_row = rows[hdr['Period']], rows[code_hdr]
    date_cols = []
    for c in range(3, len(period)):
        s = str(period[c] or '').strip()
        if re.fullmatch(r'\d{8}', s) and ITEMS.get(str(item_row[c] or '').strip()) == field:
            date_cols.append((c, datetime.date(int(s[:4]), int(s[4:6]), int(s[6:]))))
    out = []
    for r in rows[code_hdr + 1:]:
        code = str(r[0] or '').strip()
        if not code.startswith('A'):
            continue
        nm = clean_name(r[1])
        for c, d in date_cols:
            v = r[c] if c < len(r) else None
            if isinstance(v, numbers.Number) and v > 0:
                out.append((d, code, nm, float(v)))
    return pd.DataFrame(out, columns=['date', 'code', 'name', field])


def parse(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    df = None
    for sn, field in item_sheets(wb):
        g = parse_sheet(wb[sn], field)
        print(f"    {sn} · {field}: {len(g):,}행")
        df = g if df is None else df.merge(g, on=['date', 'code', 'name'], how='outer')
    if df is None:
        return pd.DataFrame(columns=['date', 'code', 'name'] + VALS)
    for v in VALS:
        if v not in df.columns:
            df[v] = None
    return df[['date', 'code', 'name'] + VALS]


def norm(d):
    d = d[['date', 'code', 'name'] + VALS].copy()
    d['date'] = d['date'].astype(str).str[:10]
    d['code'] = d['code'].astype(object)
    d['name'] = d['name'].astype(object)
    for v in VALS:
        d[v] = d[v].astype('float64')
    return d.sort_values(['date', 'code']).reset_index(drop=True)


def parse_snapshot(xlsx_path, date_key):
    """일일 스냅샷 시트(Period 'CPD')의 수정시가·고가·저가 열 -> 기준일 1일치 DataFrame.
    열이 없으면 빈 DataFrame"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hdr = {str(r[1] or '').strip(): i for i, r in enumerate(rows[:16]) if r and r[1]}
        code_hdr = next((i for i, r in enumerate(rows[:18]) if r and str(r[0] or '').strip() == 'Code'), None)
        if code_hdr is None or 'Period' not in hdr:
            continue
        period, item_row = rows[hdr['Period']], rows[code_hdr]
        cols = {ITEMS[str(item_row[c] or '').strip()]: c for c in range(3, len(item_row))
                if str(item_row[c] or '').strip() in ITEMS and str(period[c] or '').strip() == 'CPD'}
        if len(cols) < len(VALS):
            continue
        d = datetime.date.fromisoformat(date_key)
        out = []
        for r in rows[code_hdr + 1:]:
            code = str(r[0] or '').strip()
            if not code.startswith('A'):
                continue
            vals = [r[cols[v]] if cols[v] < len(r) else None for v in VALS]
            if not all(isinstance(v, numbers.Number) and v > 0 for v in vals):
                continue
            out.append((d, code, clean_name(r[1]), *[float(v) for v in vals]))
        return pd.DataFrame(out, columns=['date', 'code', 'name'] + VALS)
    return pd.DataFrame(columns=['date', 'code', 'name'] + VALS)


def build_snapshot(xlsx_path, date_key):
    """일일 파일용: 스냅샷 시트의 시가·고가·저가를 기준일 1일치로 병합"""
    df = parse_snapshot(xlsx_path, date_key)
    if df.empty:
        print('SKIP: 시가·고가·저가 열 없음')
        return False
    write(df)
    return True


def build(xlsx_path):
    """Peer Analysis 시계열 시트용 (전체 이력)"""
    df = parse(xlsx_path)
    if df.empty:
        print('SKIP: 시가·고가·저가 데이터 없음')
        return
    write(df)


def write(df):
    os.makedirs(OUT, exist_ok=True)
    con = duckdb.connect()
    n_new = n_same = 0
    for ym, g in df.groupby(df['date'].map(lambda d: f'{d.year:04d}-{d.month:02d}')):
        dst = os.path.join(OUT, ym + '.parquet')
        if os.path.exists(dst):
            old = pd.read_parquet(dst)
            old['date'] = pd.to_datetime(old['date']).dt.date
            keys = set(zip(g['date'], g['code']))
            g = pd.concat([old[[k not in keys for k in zip(old['date'], old['code'])]], g])
        g = norm(g)
        if os.path.exists(dst) and norm(pd.read_parquet(dst)).equals(g):
            n_same += 1
            continue
        con.execute(f"""
            COPY (SELECT CAST(date AS DATE) AS date, code, name,
                         CAST(open AS DOUBLE) AS open, CAST(high AS DOUBLE) AS high,
                         CAST(low AS DOUBLE) AS low
                  FROM g ORDER BY date, code)
            TO '{dst}' (FORMAT PARQUET, COMPRESSION SNAPPY)""")
        n_new += 1
    size = sum(os.path.getsize(f) for f in glob.glob(os.path.join(OUT, '*.parquet')))
    print(f"OK  OHLC -> db/market/ohlc/: {len(df):,}행, 종목 {df['code'].nunique():,}개, "
          f"{df['date'].min()} ~ {df['date'].max()} | 월 파일 {n_new}개 갱신, {n_same}개 동일 "
          f"({size // 1024 // 1024}MB)")


PRICE_DIR = os.path.join(REPO, 'db', 'market', 'price')
MIN_ROWS = 60        # 배율 판정에 필요한 최소 공통 거래일
FIT_RATIO = 0.99     # 보정 후 close가 [low, high]에 들어가야 하는 비율


def reconcile_price():
    """OHLC 기준으로 price DB의 close가 전 구간 일정 배율만큼 어긋난 종목을 찾아 소급 보정.

    OHLC 파일이 수정주가 파일보다 최신이면, 그 사이에 시행된 액면분할·병합이 고가·저가에만
    반영돼 close가 [low, high] 밖으로 나간다 (예: 5:1 병합 시 close가 정확히 1/5).
    close × r 이 (거의) 모든 날 [low, high] 안에 들어가는 배율 r을 찾으면 그것이 조정계수.
    -> 보정한 종목 수 (0이면 변경 없음)
    """
    con = duckdb.connect()
    bad = con.execute(f"""
        SELECT o.code, any_value(o.name) AS name, count(*) AS n,
               median((o.low + o.high) / 2 / p.close) AS r,
               count(*) FILTER (WHERE NOT (o.low <= p.close AND p.close <= o.high)) AS outside
        FROM '{os.path.join(OUT, '*.parquet')}' o
        JOIN '{os.path.join(PRICE_DIR, '*.parquet')}' p USING (date, code)
        GROUP BY o.code HAVING n >= {MIN_ROWS} AND outside >= n * {FIT_RATIO}
    """).df()
    if bad.empty:
        print('OK  수정주가 정합성: 배율 불일치 종목 없음')
        return 0

    fixed = []
    for _, b in bad.iterrows():
        r = round(float(b['r']), 4)
        r = round(r) if abs(r - round(r)) < 0.02 else r     # 5.0003 -> 5
        if r <= 0 or abs(r - 1) < 1e-9:
            continue
        fit = con.execute(f"""
            SELECT count(*) FILTER (WHERE o.low <= p.close * {r} AND p.close * {r} <= o.high)::DOUBLE / count(*)
            FROM '{os.path.join(OUT, '*.parquet')}' o
            JOIN '{os.path.join(PRICE_DIR, '*.parquet')}' p USING (date, code)
            WHERE o.code = '{b['code']}'
        """).fetchone()[0]
        if fit >= FIT_RATIO:
            fixed.append((b['code'], b['name'], r, int(b['n'])))
        else:
            print(f"   ! {b['code']} {b['name']}: 배율 {r} 후보이나 적합도 {fit:.1%} — 건너뜀")
    if not fixed:
        return 0

    # OHLC가 커버하는 마지막 날짜까지만 보정 (그 이후는 이미 새 기준으로 들어온 값)
    last = con.execute(f"SELECT max(date) FROM '{os.path.join(OUT, '*.parquet')}'").fetchone()[0]
    ratios = {c: r for c, _, r, _ in fixed}
    n_files = 0
    for f in sorted(glob.glob(os.path.join(PRICE_DIR, '*.parquet'))):
        df = pd.read_parquet(f)
        df['date'] = pd.to_datetime(df['date']).dt.date
        m = df['code'].isin(ratios) & (df['date'] <= last)
        if not m.any():
            continue
        df.loc[m, 'close'] = df.loc[m, 'close'] * df.loc[m, 'code'].map(ratios)
        con.execute(f"""
            COPY (SELECT CAST(date AS DATE) AS date, code, name, CAST(close AS DOUBLE) AS close
                  FROM df ORDER BY date, code)
            TO '{f}' (FORMAT PARQUET, COMPRESSION SNAPPY)""")
        n_files += 1
    print(f"OK  수정주가 소급 보정: {len(fixed)}종목 · 월 파일 {n_files}개 (~{last})")
    for c, nm, r, n in fixed:
        print(f"   {c} {nm}: close × {r} ({n:,}일)")
    return len(fixed)


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit('사용법: python3 tools/market/build_ohlc.py <수정시가·고가·저가 시트가 있는 xlsx> [--reconcile-only]')
    if '--reconcile-only' not in sys.argv[1:]:
        build(sys.argv[1])
    reconcile_price()
