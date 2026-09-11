#!/usr/bin/env python3
"""거래정지 시계열 시트(Peer Analysis 배열, 항목 '거래정지구분') -> 거래정지 DB (db/market/halt/)

원본: HTS Peer Analysis 다운로드. 행 = 종목(Code, Name, 결산월), 열 = 일자(Period 행 YYYYMMDD),
값 = '정상' / '거래정지'. 시트명과 무관하게 Code 헤더 행에 '거래정지구분' 항목이 있는 시트를 찾는다.

일일 파일(ETF_price_concensus_*.xlsx)의 스냅샷 시트('수정주가, 목표주가')에 든 기준일 1일치
'거래정지구분' 열(Period 'CPD')도 build_snapshot()으로 같은 DB에 병합한다 (ingest_daily 4단계).

산출물: db/market/halt/YYYY-MM.parquet — **거래정지인 (date, code)만** 저장
  date DATE, code VARCHAR, name VARCHAR('(주)'·'㈜' 접두 제거)
- 행이 없으면 그 날 정상 거래 (원본 전 종목이 '정상'/'거래정지' 둘 중 하나로 채워져 있음)
- 원본에 있는 날짜는 통째로 교체(그 날짜의 기존 행 삭제 후 원본의 거래정지 행 삽입),
  없는 날짜는 유지. 내용이 같은 달은 다시 쓰지 않음
- 용도: 백테스트에서 매매 불가일 제외, 신고가·RS 판정 시 정지 종목 필터 등.
  예) SELECT p.* FROM price p LEFT JOIN halt h USING (date, code) WHERE h.code IS NULL
실행: python3 tools/market/build_halt.py <파일.xlsx>   (ingest_daily.py가 시트 구조로 자동 인식)
"""
import sys, os, re, datetime
import openpyxl
import duckdb
import pandas as pd

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(REPO, 'db', 'market', 'halt')
ITEM = '거래정지구분'
HALT = '거래정지'


def clean_name(s):
    return re.sub(r'^\(주\)|^㈜', '', str(s or '').strip()).strip()


def find_sheet(wb):
    """Code 헤더 행에 '거래정지구분' 항목이 있는 시트명 (없으면 None)"""
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 16:
                break
            if row and str(row[0] or '').strip() == 'Code':
                if any(str(v or '').strip() == ITEM for v in row[3:8]):
                    return ws.title
                break
    return None


def is_halt_file(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    return find_sheet(wb) is not None


def parse(xlsx_path):
    """-> (DataFrame(date, code, name) 거래정지 행, 원본이 담은 날짜 목록)"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    sn = find_sheet(wb)
    if not sn:
        return pd.DataFrame(columns=['date', 'code', 'name']), []
    rows = list(wb[sn].iter_rows(values_only=True))
    hdr = {str(r[1] or '').strip(): i for i, r in enumerate(rows[:16]) if r and r[1]}
    period = rows[hdr['Period']]
    code_hdr = next(i for i, r in enumerate(rows[:18]) if r and str(r[0] or '').strip() == 'Code')
    item_row = rows[code_hdr]
    date_cols = []
    for c in range(3, len(period)):
        s = str(period[c] or '').strip()
        if re.fullmatch(r'\d{8}', s) and str(item_row[c] or '').strip() == ITEM:
            date_cols.append((c, datetime.date(int(s[:4]), int(s[4:6]), int(s[6:]))))
    out = []
    for r in rows[code_hdr + 1:]:
        code = str(r[0] or '').strip()
        if not code.startswith('A'):
            continue
        nm = clean_name(r[1])
        for c, d in date_cols:
            v = r[c] if c < len(r) else None
            if v is not None and str(v).strip() == HALT:
                out.append((d, code, nm))
    return pd.DataFrame(out, columns=['date', 'code', 'name']), [d for _, d in date_cols]


def parse_snapshot(xlsx_path, date_key):
    """일일 스냅샷 시트(기준일 1일치, Period 'CPD')의 '거래정지구분' 열 -> (거래정지 행 DataFrame, [기준일])
    열이 없으면 ([], [])"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    sn = find_sheet(wb)
    if not sn:
        return pd.DataFrame(columns=['date', 'code', 'name']), []
    rows = list(wb[sn].iter_rows(values_only=True))
    hdr = {str(r[1] or '').strip(): i for i, r in enumerate(rows[:16]) if r and r[1]}
    code_hdr = next(i for i, r in enumerate(rows[:18]) if r and str(r[0] or '').strip() == 'Code')
    period, item_row = rows[hdr['Period']], rows[code_hdr]
    col = next((c for c in range(3, len(item_row))
                if str(item_row[c] or '').strip() == ITEM and str(period[c] or '').strip() == 'CPD'), None)
    if col is None:
        return pd.DataFrame(columns=['date', 'code', 'name']), []
    d = datetime.date.fromisoformat(date_key)
    out = [(d, str(r[0]).strip(), clean_name(r[1])) for r in rows[code_hdr + 1:]
           if r and str(r[0] or '').strip().startswith('A') and col < len(r)
           and r[col] is not None and str(r[col]).strip() == HALT]
    return pd.DataFrame(out, columns=['date', 'code', 'name']), [d]


def build_snapshot(xlsx_path, date_key):
    """일일 파일용: 스냅샷 시트의 거래정지구분 열을 기준일 1일치로 병합"""
    df, dates = parse_snapshot(xlsx_path, date_key)
    if not dates:
        print('SKIP: 거래정지구분 열 없음')
        return
    write(df, dates)


def build(xlsx_path):
    """Peer Analysis 시계열 시트용"""
    df, dates = parse(xlsx_path)
    if not dates:
        print('SKIP: 거래정지 시트 없음')
        return
    write(df, dates)


def write(df, dates):
    os.makedirs(OUT, exist_ok=True)
    con = duckdb.connect()
    n_new = n_same = 0
    by_month = {}
    for d in dates:
        by_month.setdefault(f'{d.year:04d}-{d.month:02d}', set()).add(d)
    for ym, dset in sorted(by_month.items()):
        dst = os.path.join(OUT, ym + '.parquet')
        g = df[df['date'].isin(dset)]
        if os.path.exists(dst):
            old = pd.read_parquet(dst)
            old['date'] = pd.to_datetime(old['date']).dt.date
            g = pd.concat([old[~old['date'].isin(dset)], g], ignore_index=True)
        g = g.sort_values(['date', 'code']).reset_index(drop=True)
        g['date'] = g['date'].astype(str).str[:10]
        if os.path.exists(dst):
            o = pd.read_parquet(dst)
            o['date'] = o['date'].astype(str).str[:10]
            if o.sort_values(['date', 'code']).reset_index(drop=True).equals(g):
                n_same += 1
                continue
        con.execute(f"""
            COPY (SELECT CAST(date AS DATE) AS date, code, name FROM g ORDER BY date, code)
            TO '{dst}' (FORMAT PARQUET, COMPRESSION SNAPPY)""")
        n_new += 1
    print(f"OK  거래정지 -> db/market/halt/: {len(df):,}행(정지일×종목), 종목 {df['code'].nunique():,}개, "
          f"원본 {min(dates)} ~ {max(dates)} {len(dates)}일 | 월 파일 {n_new}개 갱신, {n_same}개 동일")


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit('사용법: python3 tools/market/build_halt.py <거래정지 시트가 있는 xlsx>')
    build(sys.argv[1])
