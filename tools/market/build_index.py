#!/usr/bin/env python3
"""kospi_kosdaq.xlsx -> 시장지수 DB (db/market/index/)

원본: HTS 'Time Series (Sector)' 다운로드 (1시트). 헤더 행에 Code(IKS900 코스피 /
IKQ900 코스닥)·Name·Item Code, 'D A T E' 행 아래로 일별 종가·시가·고가·저가 지수.
1999-12-28부터의 전체 이력이 담기며, 마지막 행(당일 CPD)은 값이 비어 있어 제외된다.

산출물: db/market/index/YYYY.parquet (연도별)
  date DATE, code VARCHAR(IKS900/IKQ900), name VARCHAR(코스피/코스닥),
  open/high/low/close DOUBLE (지수 포인트)
- 원본에 있는 날짜만 교체하고 나머지는 유지(병합). 내용이 같은 연도 파일은 다시 쓰지 않음
- 백테스트에서 벤치마크·시장 국면 판단용. 예: duckdb.sql("SELECT * FROM 'db/market/index/*.parquet'")
실행: python3 tools/market/build_index.py <kospi_kosdaq.xlsx>
      (ingest_daily.py가 시트 구조로 자동 인식해 호출)
"""
import sys, os, numbers, datetime
import openpyxl
import duckdb
import pandas as pd

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(REPO, 'db', 'market', 'index')

ITEMS = {'종가지수': 'close', '시가지수': 'open', '고가지수': 'high', '저가지수': 'low'}
COLS = ['date', 'code', 'name', 'open', 'high', 'low', 'close']


def is_index_file(xlsx_path):
    """첫 시트 Code 행이 지수 코드(IKS/IKQ…)로 시작하면 지수 파일"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 12:
            break
        if row and str(row[0] or '').strip() == 'Code':
            return any(str(v or '').startswith('IK') for v in row[1:])
    return False


def parse(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = {str(r[0]).strip(): i for i, r in enumerate(rows[:16]) if r and r[0]}
    code_row, name_row, item_row = rows[header['Code']], rows[header['Name']], rows[header['D A T E']]
    cols = []   # (col_idx, code, name, field)
    for c in range(1, len(code_row)):
        code = str(code_row[c] or '').strip()
        field = ITEMS.get(str(item_row[c] or '').strip())
        if code and field:
            cols.append((c, code, str(name_row[c] or '').strip(), field))

    recs = {}   # (date, code) -> {name, open, high, low, close}
    for r in rows[header['D A T E'] + 1:]:
        d = r[0] if r else None
        if isinstance(d, datetime.datetime):
            d = d.date()
        elif d is not None:
            try:
                d = datetime.date.fromisoformat(str(d)[:10])
            except ValueError:
                continue
        else:
            continue
        for c, code, name, field in cols:
            v = r[c] if c < len(r) else None
            if isinstance(v, numbers.Number):
                recs.setdefault((d, code), {'name': name})[field] = float(v)
    out = [(d, code, x['name'], x.get('open'), x.get('high'), x.get('low'), x.get('close'))
           for (d, code), x in recs.items() if x.get('close') is not None]
    return pd.DataFrame(out, columns=COLS)


def norm(d):
    d = d.copy()
    d['date'] = d['date'].astype(str).str[:10]
    for k in ('open', 'high', 'low', 'close'):
        d[k] = d[k].astype('float64')
    return d.sort_values(['date', 'code']).reset_index(drop=True)


def build(xlsx_path):
    df = parse(xlsx_path)
    if df.empty:
        print('SKIP: 지수 데이터 없음')
        return
    os.makedirs(OUT, exist_ok=True)
    n_new = n_same = 0
    con = duckdb.connect()
    for y, g in df.groupby(df['date'].map(lambda d: d.year)):
        dst = os.path.join(OUT, f'{y}.parquet')
        if os.path.exists(dst):
            old = pd.read_parquet(dst)
            old['date'] = pd.to_datetime(old['date']).dt.date
            keys = set(zip(g['date'], g['code']))
            g = pd.concat([old[[k not in keys for k in zip(old['date'], old['code'])]], g])
        g = norm(g)
        if os.path.exists(dst) and norm(pd.read_parquet(dst)).equals(g):
            n_same += 1
            continue
        con.execute(f"""
            COPY (SELECT CAST(date AS DATE) AS date, code, name,
                         CAST(open AS DOUBLE) AS open, CAST(high AS DOUBLE) AS high,
                         CAST(low AS DOUBLE) AS low, CAST(close AS DOUBLE) AS close
                  FROM g ORDER BY date, code)
            TO '{dst}' (FORMAT PARQUET, COMPRESSION SNAPPY)""")
        n_new += 1
    names = ', '.join(f"{n}({c})" for c, n in df.groupby('code')['name'].first().items())
    print(f"OK  시장지수 -> db/market/index/: {len(df):,}행 ({df['date'].min()} ~ {df['date'].max()}) "
          f"| {names} | 연도 파일 {n_new}개 갱신, {n_same}개 동일")


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit('사용법: python3 tools/market/build_index.py <kospi_kosdaq.xlsx>')
    build(sys.argv[1])
